"""Regression for post-training Excel export of large diagnostic/config fields.

Runs with unittest as well as pytest. No model construction or PyTorch needed.
"""
import csv
from contextlib import closing
import hashlib
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import load_workbook

from v52.export import _large_csv_fields,export_workbook


class LargeCSVExportTests(unittest.TestCase):
    def setUp(self):
        self.previous=csv.field_size_limit()
        csv.field_size_limit(131072)
        self.temp=tempfile.TemporaryDirectory()
        self.report=Path(self.temp.name)/'publication'
        self.tables=self.report/'tables';self.tables.mkdir(parents=True)

    def tearDown(self):
        csv.field_size_limit(self.previous)
        self.temp.cleanup()

    def write_csv(self,name,rows):
        path=self.tables/name
        with path.open('w',newline='',encoding='utf8') as stream:
            csv.writer(stream).writerows(rows)
        return path

    def test_report_reads_field_larger_than_default_and_keeps_original(self):
        payload=json.dumps({'trace':['context-radius update']*15000},ensure_ascii=False)
        self.assertGreater(len(payload),131072)
        source=self.write_csv('phase_and_execution_diagnostics.csv',[
            ['method','details_json'],['Method3',payload],['Method7','small diagnostic']])
        before=source.read_bytes()
        # This is the user's failure before export's oversized-cell branch runs.
        with source.open(newline='',encoding='utf8') as stream:
            reader=csv.reader(stream);next(reader)
            with self.assertRaisesRegex(csv.Error,'field larger than field limit'):
                next(reader)
        manifest=export_workbook(self.report)
        self.assertEqual(source.read_bytes(),before)
        self.assertEqual(manifest['sources'][0]['sha256'],hashlib.sha256(before).hexdigest())
        self.assertEqual(manifest['sources'][0]['rows'],2)
        with closing(load_workbook(self.report/manifest['workbook'],read_only=True)) as book:
            sheet=book[manifest['sources'][0]['sheets'][0]]
            self.assertEqual(sheet['B2'].value,'[Full value in tables/phase_and_execution_diagnostics.csv; CSV row 2]')
            self.assertEqual(sheet['B3'].value,'small diagnostic')
        self.assertEqual(csv.field_size_limit(),131072)
        self.assertFalse((self.report/'.research_tables.partial.xlsx').exists())

    def test_excel_cell_boundary_and_csv_limit_boundary(self):
        sizes=[32767,32768,131072,131073,200001]
        source=self.write_csv('configurations.csv',[
            ['length','config_key'],*[[length,'x'*length] for length in sizes]])
        before=source.read_bytes();manifest=export_workbook(self.report)
        with closing(load_workbook(self.report/manifest['workbook'],read_only=True)) as book:
            sheet=book[manifest['sources'][0]['sheets'][0]]
            for number,length in enumerate(sizes,2):
                with self.subTest(length=length):
                    expected='x'*length if length<=32767 else f'[Full value in tables/configurations.csv; CSV row {number}]'
                    self.assertEqual(sheet.cell(number,2).value,expected)
        self.assertEqual(source.read_bytes(),before)
        self.assertEqual(manifest['sources'][0]['rows'],len(sizes))
        self.assertEqual(csv.field_size_limit(),131072)

    def test_unicode_quoted_multiline_diagnostics_and_other_values(self):
        large=('فارسی، "quoted"\n'*12000)
        source=self.write_csv('diagnostics.csv',[
            ['details','pair_id','rmse','literal'],
            [large,'9007199254740993','0.8767','=1+1'],
            ['فارسی، "quoted"\nline two','42','0.9','https://example.com']])
        before=source.read_bytes();manifest=export_workbook(self.report)
        with closing(load_workbook(self.report/manifest['workbook'],read_only=True)) as book:
            sheet=book[manifest['sources'][0]['sheets'][0]]
            self.assertEqual(sheet['A2'].value,'[Full value in tables/diagnostics.csv; CSV row 2]')
            self.assertEqual(sheet['B2'].value,'9007199254740993')
            self.assertEqual(sheet['C2'].value,0.8767)
            self.assertEqual(sheet['D2'].value,'=1+1');self.assertEqual(sheet['D2'].data_type,'s')
            self.assertEqual(sheet['A3'].value,'فارسی، "quoted"\nline two')
            self.assertEqual(sheet['D3'].value,'https://example.com')
        self.assertEqual(source.read_bytes(),before)
        self.assertEqual(csv.field_size_limit(),131072)

    def test_limit_is_portable_to_windows_c_long(self):
        real_limit=csv.field_size_limit;requested=[]
        def windows_limit(value=None):
            if value is None:return real_limit()
            if value>2**31-1:raise OverflowError('Python int too large to convert to C long')
            requested.append(value);return real_limit(value)
        self.write_csv('diagnostics.csv',[['details'],['x'*140000]])
        with patch('v52.export.csv.field_size_limit',side_effect=windows_limit):
            export_workbook(self.report)
        self.assertIn(2**31-1,requested)
        self.assertEqual(real_limit(),131072)

    def test_larger_existing_limit_is_not_lowered(self):
        state={'limit':2**31+17}
        def field_limit(value=None):
            before=state['limit']
            if value is not None:state['limit']=value
            return before
        with patch('v52.export.csv.field_size_limit',side_effect=field_limit):
            with _large_csv_fields():self.assertEqual(state['limit'],2**31+17)
        self.assertEqual(state['limit'],2**31+17)

    def test_error_restores_limit_and_preserves_previous_workbook(self):
        self.write_csv('metrics.csv',[['method','rmse'],['Method1','.9']])
        export_workbook(self.report)
        target=self.report/'research_tables.xlsx';before=target.read_bytes()
        self.write_csv('too_many_columns.csv',[[str(i) for i in range(16385)]])
        with self.assertRaisesRegex(ValueError,'Excel column limit'):
            export_workbook(self.report)
        self.assertEqual(csv.field_size_limit(),131072)
        self.assertEqual(target.read_bytes(),before)
        self.assertFalse((self.report/'.research_tables.partial.xlsx').exists())

    def test_empty_csv_restores_limit(self):
        self.write_csv('empty.csv',[])
        manifest=export_workbook(self.report)
        self.assertEqual(csv.field_size_limit(),131072)
        self.assertEqual(manifest['sources'],[])
        self.assertTrue((self.report/manifest['workbook']).is_file())


if __name__=='__main__':unittest.main()
